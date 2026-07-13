"""Create the original, synthetic DOCX fixture for Spike 001.

This generator is test infrastructure only. It is not the clean-room DOCX
editing tool being evaluated by the spike.
"""

from pathlib import Path
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


FIXTURE_CANARY = "SPIKE001-ORIGINAL-CANARY-7F93D1"


def set_style_font(style, name: str, size: int, color: str | None = None) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style._element.rPr.rFonts.set(qn("w:ascii"), name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if color:
        style.font.color.rgb = RGBColor.from_string(color)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    set_style_font(normal, "Calibri", 11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_1 = document.styles["Heading 1"]
    set_style_font(heading_1, "Calibri", 16, "2E74B5")
    heading_1.paragraph_format.space_before = Pt(16)
    heading_1.paragraph_format.space_after = Pt(8)

    heading_2 = document.styles["Heading 2"]
    set_style_font(heading_2, "Calibri", 13, "2E74B5")
    heading_2.paragraph_format.space_before = Pt(12)
    heading_2.paragraph_format.space_after = Pt(6)


def build_fixture(output_path: Path) -> None:
    document = Document()
    configure_document(document)

    document.add_heading("Spike 001 Synthetic DOCX Fixture", level=1)
    document.add_paragraph(
        "This disposable document exists only to validate a safe, revised-copy "
        "DOCX workflow. It contains no user or proprietary source material."
    )

    document.add_heading("Executive Summary", level=1)
    document.add_paragraph(
        "The Northstar pilot demonstrated that a focused document workflow can "
        "reduce repetitive review effort while preserving a clear human approval "
        "boundary. The initial evaluation covered a deliberately narrow scenario "
        "so that file safety, sandbox enforcement, and output validation could be "
        "measured independently before broader product work begins."
    )
    document.add_paragraph(
        "The recommended next step is to validate the complete execution boundary "
        "with synthetic inputs, record every assumption, and proceed only when the "
        "original document remains unchanged and all restricted operations fail "
        "closed."
    )

    document.add_heading("Operating Constraints", level=1)
    document.add_paragraph(
        "This adjacent section must remain unchanged during the Executive Summary "
        "rewrite. Network access is not required by the document tool."
    )

    document.add_heading("Validation Canary", level=2)
    document.add_paragraph(FIXTURE_CANARY)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)

    reopened = Document(output_path)
    full_text = "\n".join(p.text for p in reopened.paragraphs)
    required = ["Executive Summary", "Operating Constraints", FIXTURE_CANARY]
    missing = [value for value in required if value not in full_text]
    if missing:
        raise RuntimeError(f"Fixture validation failed; missing: {missing}")


def main() -> int:
    if len(sys.argv) != 2:
        print("usage: create_synthetic_docx_fixture.py <output.docx>", file=sys.stderr)
        return 2
    output_path = Path(sys.argv[1]).resolve()
    if output_path.suffix.lower() != ".docx":
        print("output must end in .docx", file=sys.stderr)
        return 2
    build_fixture(output_path)
    print(output_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
